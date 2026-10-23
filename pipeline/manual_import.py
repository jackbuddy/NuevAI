"""Manual-documents stage: files you download yourself and drop in a folder.

For Google Docs that are private (or PDFs / Drive files) the crawler cannot
read, open them in your browser, File -> Download (.docx / .pdf / .txt), and put
the files in  data/pipeline/manual_docs/ . On the next run they are indexed.
The file name (without extension) becomes the document title, so name them
like people would ask for them ("Student Forms.docx").

Supported: .txt .md .html .csv .docx .xlsx .pdf
"""
import re
from datetime import datetime, timezone

from bs4 import BeautifulSoup

from . import config, workbooks
from .common import html_to_markdown, save_page

MANUAL_DIR = config.DATA_DIR / "manual_docs"

# Spreadsheet tabs that hold individual students (advisee lists, rosters, name lists) are never indexed, even if they
# are left in the export. Skipped tab names are reported in the run summary so nothing disappears silently.
SKIP_TAB_TITLE = re.compile(r"advisee|roster|student (list|names?)|class list|attendance|grades?\b|medical|diploma", re.I)
SKIP_TAB_HEADER = re.compile(r"advisee|student\s*(last|first|name|id)|^students?$", re.I)
SKIPPED_TABS: list[str] = []
# text files the old scraper already pulled by hand
LEGACY_DIR = config.NEXUS_SCRAPER_DIR / "data" / "manual_docs"


def read_file(path) -> str:
    ext = path.suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext in (".html", ".htm"):
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        for s in soup.select("script, style"):
            s.decompose()
        return html_to_markdown(soup.body or soup)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        out = []
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            m = re.match(r"Heading (\d)", p.style.name or "")
            out.append(f"{'#' * int(m.group(1))} {t}" if m else t)
        for table in d.tables:
            for row in table.rows:
                out.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        return "\n\n".join(out)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:                       # every tab, as comma-separated rows under a heading
            rows = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            head = rows[:3]
            if SKIP_TAB_TITLE.search(ws.title) or any(SKIP_TAB_HEADER.search(c) for r in head for c in r if c):
                SKIPPED_TABS.append(f"{path.name} > {ws.title}")
                continue
            out.append(f"## {ws.title}")
            out.extend(",".join(r) for r in rows)
        return "\n".join(out)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
    raise ValueError(f"unsupported file type {ext}")


def run(log=print) -> dict:
    MANUAL_DIR.mkdir(parents=True, exist_ok=True)
    imported, failed, handled = 0, [], []
    files = sorted(MANUAL_DIR.glob("*")) + (sorted(LEGACY_DIR.glob("*.txt")) if LEGACY_DIR.exists() else [])
    for path in files:
        if path.name.startswith(".") or path.is_dir():
            continue
        if path.suffix.lower() == ".xlsx" and workbooks.kind_of_path(path):
            handled.append(path.name)                  # parsed into structured records by pipeline/workbooks.py
            continue
        try:
            text = read_file(path).strip()
        except Exception as e:
            failed.append(f"{path.name}: {e}")
            continue
        if len(text) < 80:
            failed.append(f"{path.name}: no readable text (scanned PDF?)")
            continue
        legacy = path.parent == LEGACY_DIR
        first = next((l.lstrip("# ").strip() for l in text.splitlines() if l.strip()), path.stem)
        title = (first[:100] if legacy else path.stem.replace("_", " ").strip())
        save_page({"url": f"manual://{path.name}", "title": title, "source": "manual", "text": text, "links": [],
                   "scraped_at": datetime.now(timezone.utc).isoformat()})
        imported += 1
    stats = {"imported": imported, "failed": failed, "structured_workbooks": handled, "skipped_private_tabs": SKIPPED_TABS, "folder": str(MANUAL_DIR)}
    log(f"manual docs done: {stats}")
    return stats
