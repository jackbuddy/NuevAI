"""Structured parsers for the school's planning workbooks (Excel exports or live Google Sheets).

Grids that people read by eye (a week calendar, teacher-by-block tables) turn into nonsense when flattened to text, so
each known workbook layout is parsed into explicit records:

  master schedule   "The Sheet": A/B day + events for every date, advisory times by grade, teacher -> course -> block
                    -> room (Fall/Spring "Roles"), advisor rooms. Student-name tabs are never read.
  course offerings  every course by block with the teacher's last name; merged with the master schedule's full names.
  clubs catalog     one record per club (leaders, advisor, day, time, room).

Layouts are detected from content, not file names. Anything unrecognised is left to the generic importer.
"""
import hashlib
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

from . import config

BLOCK_RULE = ("Blocks 1-4 meet on A and A' days; Blocks 5-8 meet on B and B' days. A' and B' flip the last two blocks "
              "(A' order 1, 2, 4, 3; B' order 5, 6, 8, 7).")
SLOT_TIMES = ["8:55-10:10 AM", "10:15-11:30 AM", "12:35-1:50 PM", "1:55-3:10 PM"]     # confirmed bell times (lunch 11:30-12:00/12:30)
ORDER = {"A": [1, 2, 3, 4], "A'": [1, 2, 4, 3], "B": [5, 6, 7, 8], "B'": [5, 6, 8, 7]}
# spellings the school confirmed: the Roles tab has a typo
NAME_FIXES = {"Jeremy Jaquot": "Jeremy Jacquot"}
DAY_TYPES = {"A", "A'", "B", "B'"}
STUDENT_TABS = re.compile(r"roster|group(s)?$|advisee|student", re.I)        # never parsed


# --------------------------------------------------------------------------- helpers
def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return re.sub(r"[ \t]+", " ", str(v).replace("\xa0", " ")).strip()


def grid(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s.lower().replace("&", " and ").replace("+", " and "))).strip()


def quote(s: str) -> str:
    return s.replace("’", "'").replace("′", "'").strip()


def mk(seed: str, file: str, title: str, section: str, body: str, meta: dict) -> dict:
    content = f"Title: {title}\nSection: {section}\nSource: {file}\n\n{body}"
    cid = "wb-" + hashlib.sha1(seed.encode()).hexdigest()[:14]
    return {"chunk_id": cid, "content": content, "source_file": f"workbook://{file}"[:250],
            "metadata": {"source_url": f"workbook://{file}", "title": title, **meta,
                         "content_hash": hashlib.sha1((content + repr(sorted(meta.items()))).encode()).hexdigest()}}


def kind_of(wb) -> str | None:
    names = {ws.title.strip() for ws in wb.worksheets}
    if {"Calendar AB + Special Schedule", "Fall Roles"} <= names:
        return "master"
    first = next((ws for ws in wb.worksheets), None)
    if first is not None:
        head = [cell(c).lower() for c in next(iter(grid(first)), [])]
        if head[:3] == ["clubs", "description", "leaders"]:
            return "clubs"
    for ws in wb.worksheets:
        top = " ".join(cell(c) for r in grid(ws)[:3] for c in r)
        if "Block 1" in top and "Block 8" in top and "Roles" not in ws.title and "Rooms" not in ws.title:
            return "offerings"
    return None


def kind_of_path(path: Path) -> str | None:
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception:
        return None
    return kind_of(wb)


# --------------------------------------------------------------------------- master schedule: days
def parse_days(wb) -> dict[date, dict]:
    ws = wb["Calendar AB + Special Schedule"]
    rows = grid(ws)
    days: dict[date, dict] = {}
    ym = None
    i = 0
    while i < len(rows):
        r = rows[i] + [None] * 8
        marker = r[1]
        if isinstance(marker, datetime) and not any(cell(r[j]) for j in (2, 3, 4, 5)):
            ym = (marker.year, marker.month)
            i += 1
            continue
        nums = [(j, r[j]) for j in range(1, 6) if isinstance(r[j], (int, float)) and 1 <= r[j] <= 31]
        if nums and ym and i + 2 < len(rows):
            ev, ty = rows[i + 1] + [None] * 8, rows[i + 2] + [None] * 8
            for j, d in nums:
                try:
                    dd = date(ym[0], ym[1], int(d))
                except ValueError:
                    continue
                if dd.weekday() != j - 1:          # sanity check: column must match the weekday
                    continue
                label, kind = cell(ev[j]), quote(cell(ty[j]))
                days[dd] = {"labels": [x.strip() for x in re.split(r"\s*/\s*|\n", label) if x.strip()], "type": kind}
            i += 3
            continue
        i += 1
    return days


def parse_advisory(wb) -> dict[date, dict[str, str]]:
    """date -> {grade: what happens that day}. Plain 'Advisory' gets the grade's usual time window."""
    if "Advisory Calendar" not in wb.sheetnames:
        return {}
    rows = grid(wb["Advisory Calendar"])
    grades = ["Freshmen", "Sophomores", "Juniors", "Seniors"]
    usual = {g: cell(rows[1][3 + k]) if len(rows) > 1 and len(rows[1]) > 3 + k else "" for k, g in enumerate(grades)}
    out: dict[date, dict[str, str]] = {}
    for r in rows[2:]:
        r = r + [None] * 8
        d = r[1]
        if not isinstance(d, datetime):
            continue
        entry = {}
        for k, g in enumerate(grades):
            v = re.sub(r"\s*\n\s*", " ", cell(r[3 + k]))
            if v:
                entry[g] = f"{v} ({usual[g]})" if v.lower() == "advisory" and usual[g] else v
        if entry:
            out[d.date()] = entry
    return out


def day_chunks(wb, file: str) -> list[dict]:
    days, adv = parse_days(wb), parse_advisory(wb)
    chunks = []
    for d in sorted(set(days) | set(adv)):
        info = days.get(d, {})
        label = ", ".join(info.get("labels", []))
        t = info.get("type", "")
        if t in DAY_TYPES:
            order = ORDER[t]
            slots = [f"Block {order[0]} {SLOT_TIMES[0]}", f"Block {order[1]} {SLOT_TIMES[1]}", "Lunch 11:30 AM-12:00/12:30 PM",
                     f"Block {order[2]} {SLOT_TIMES[2]}", f"Block {order[3]} {SLOT_TIMES[3]}"]
            lead = f"{t} day: Blocks {order[0]}-{max(order)} meet, in the order {', '.join(map(str, order))}. Times: {'; '.join(slots)}."
            status = lead + (f" Schedule notes: {label} (special events can adjust the regular times)." if label else "")
            kind = "class_day"
        elif t or "no class" in label.lower():
            status = f"No classes{(' (' + t + ')') if t else ''}." + (f" {label}." if label and label.lower() != "no classes" else "")
            kind = "no_classes"
        else:
            status = f"{label}." if label else ""
            kind = "other"
        parts = [status] if status else []
        a = adv.get(d)
        if a:
            vals = set(a.values())
            parts.append("Advisory: " + (f"all grades - {next(iter(vals))}" if len(vals) == 1 and len(a) == 4
                                          else "; ".join(f"{g}: {v}" for g, v in a.items())) + ".")
        if not parts:
            continue
        head = d.strftime("%A, %B %-d, %Y")
        body = f"{head} - " + " ".join(parts)
        if kind == "class_day":
            body += f" ({BLOCK_RULE})"
        chunks.append(mk(f"day|{d}", file, f"School day type: {head}", "A/B day calendar and advisory", body,
                         {"source": "school_day", "kind": kind, "event_date": d.isoformat(), "end_date": d.isoformat(),
                          "day_type": t if t in DAY_TYPES else ""}))
    return chunks


# --------------------------------------------------------------------------- master schedule: advisors
def advisor_chunks(wb, file: str) -> list[dict]:
    chunks = []
    if "Advisors + Locations" in wb.sheetnames:
        rows = grid(wb["Advisors + Locations"])
        teams: dict[str, list[str]] = defaultdict(list)
        team = ""
        for r in rows:
            r = r + [None] * 6
            c0 = cell(r[0])
            if c0 and not cell(r[1]) and re.search(r"advis\w* team", c0, re.I):
                team = c0
                continue
            num = cell(r[2])
            if team and re.fullmatch(r"\d+\.\d+", num):
                names = [x for x in (c0, cell(r[1])) if x]
                room = cell(r[3])
                teams[team].append(f"- Advisory {num}: {' and '.join(names)}" + (f", room {room}" if room else ""))
        for team, lines in teams.items():
            chunks.append(mk(f"adv|{team}", file, f"{team}: advisors and rooms", "Advisors + Locations",
                             f"{team} (regular advisory rooms, 2026-27):\n" + "\n".join(lines),
                             {"source": "school_advisory", "kind": "advisors"}))
    if "Intergrade Advisors" in wb.sheetnames:
        lines = []
        for r in grid(wb["Intergrade Advisors"])[1:]:
            r = r + [None] * 6
            if cell(r[0]):
                advs = ", ".join(cell(x) for x in r[2:6] if cell(x))
                lines.append(f"- Group {cell(r[0])}: meets at {cell(r[1])}; advisors: {advs}")
        if lines:
            chunks.append(mk("intergrade", file, "Intergrade advisory groups: locations and advisors", "Intergrade Advisors",
                             "Intergrade groups (advisors and where each group meets):\n" + "\n".join(lines),
                             {"source": "school_advisory", "kind": "intergrade"}))
    if "BTSN Advisory Locations" in wb.sheetnames:
        lines, team = [], ""
        for r in grid(wb["BTSN Advisory Locations"]):
            r = r + [None] * 5
            c0 = cell(r[0])
            if c0 and not cell(r[1]) and re.search(r"advis\w* team", c0, re.I):
                team = c0
            elif c0 and re.fullmatch(r"\d+\.\d+", cell(r[2])):
                names = " and ".join(x for x in (c0, cell(r[1])) if x)
                lines.append(f"- {team or 'Advisory'}: advisory {cell(r[2])}, {names}, at {cell(r[3])}")
        if lines:
            chunks.append(mk("btsn", file, "Back to School Night: where each advisory meets", "BTSN Advisory Locations",
                             "Back to School Night advisory locations:\n" + "\n".join(lines),
                             {"source": "school_advisory", "kind": "btsn"}))
    return chunks


# --------------------------------------------------------------------------- courses and teachers
def parse_roles(wb, tab: str, semester: str, known_last: frozenset = frozenset(), extra_names: set | None = None) -> list[dict]:
    """Rows of (teacher, course, block, room, dept, home_room) from a '<Fall|Spring> Roles' tab."""
    ws = wb[tab]
    rows = [r + [None] * 14 for r in grid(ws)]
    hdr = next((i for i, r in enumerate(rows[:6]) if "Block 1" in [cell(c) for c in r]), None)
    if hdr is None:
        return []
    cols = {}
    for j, c in enumerate(rows[hdr]):
        m = re.fullmatch(r"Block (\d)", cell(c))
        if m:
            cols[int(m.group(1))] = j
    room_col = next((j for j, c in enumerate(rows[hdr - 1]) if cell(c).lower() == "room"), None)
    out, dept = [], ""
    for r in rows[hdr + 1:]:
        name = NAME_FIXES.get(cell(r[0]), cell(r[0]))
        blocks = {b: cell(r[j]) for b, j in cols.items() if cell(r[j]) not in ("", "/")}
        if not name or name.lower() in ("type", "room"):
            continue
        if any(re.search(r"\bDays$", v) for v in blocks.values()):      # a repeated column-header row, not a teacher
            continue
        if not blocks:
            # a teacher with no assignments entered yet looks just like a department heading; the offerings sheet's
            # last names tell them apart
            if extra_names is not None and last_name(name) in known_last and len(name.split()) >= 2:
                extra_names.add(name)
            else:
                dept = name
            continue
        home = cell(r[room_col]) if room_col is not None else ""
        for b, text in blocks.items():
            lines = [l.strip() for l in re.split(r"\n", str(r[cols[b]])) if l.strip()]
            room = ""
            if len(lines) > 1 and len(cell(lines[-1])) <= 14:
                room, lines = cell(lines[-1]), lines[:-1]
            course = cell(" ".join(lines))
            if course:
                out.append({"semester": semester, "teacher": name, "course": course, "block": b, "room": room,
                            "dept": dept, "home": home})
    return out


def parse_offerings(wb) -> list[dict]:
    """(course, block, teacher_last) from the offerings sheet's 'Sheet1' grid ('Course\\nTeacher' cells)."""
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    rows = grid(ws)
    hdr = next((i for i, r in enumerate(rows[:4]) if "Block 1" in [cell(c) for c in r]), 0)
    cols = {int(m.group(1)): j for j, c in enumerate(rows[hdr]) if (m := re.fullmatch(r"Block (\d)", cell(c)))}
    out = []
    for r in rows[hdr + 1:]:
        r = r + [None] * 12
        for b, j in cols.items():
            lines = [l.strip() for l in str(r[j] or "").split("\n") if l.strip()]
            if not lines:
                continue
            course = cell(lines[0])
            teacher = cell(lines[1]) if len(lines) > 1 else ""
            if course and course != "/":
                out.append({"semester": "Fall 2026", "course": course, "block": b, "teacher_last": teacher})
    return out


def edit1(a: str, b: str) -> bool:
    """True when two words differ by at most one insertion, deletion or substitution (same first letter, len >= 5)."""
    if a == b:
        return True
    if len(a) < 5 or len(b) < 5 or a[0] != b[0] or abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) <= 1
    s, l = (a, b) if len(a) < len(b) else (b, a)
    i = j = diff = 0
    while i < len(s) and j < len(l):
        if s[i] == l[j]:
            i += 1
            j += 1
        else:
            diff += 1
            j += 1
            if diff > 1:
                return False
    return True


def last_name(full: str) -> str:
    return full.strip().split(" ")[-1].lower() if full else ""


def course_chunks(master_wb, off_wb, file_master: str, file_off: str) -> list[dict]:
    recs = []                                  # unified assignments
    off_recs = parse_offerings(off_wb) if off_wb is not None else []
    known_last = frozenset(r["teacher_last"].lower().strip() for r in off_recs if r["teacher_last"])
    extra_names: set[str] = set()
    if master_wb is not None:
        for tab, sem in (("Fall Roles", "Fall 2026"), ("Spring Roles", "Spring 2027 (draft: still being updated)")):
            if tab in master_wb.sheetnames:
                recs += [{**r, "src": "roles"} for r in parse_roles(master_wb, tab, sem, known_last, extra_names)]
    recs += [{**r, "src": "offerings", "teacher": "", "room": "", "dept": "", "home": ""} for r in off_recs]
    full_by_last: dict[str, set] = defaultdict(set)
    for n in extra_names:
        full_by_last[last_name(n)].add(n)
    for r in recs:
        if r.get("teacher"):
            full_by_last[last_name(r["teacher"])].add(r["teacher"])
    aliases: dict[str, set] = defaultdict(set)         # full name -> other spellings of the last name seen elsewhere
    merged: dict[tuple, dict] = {}
    for r in recs:
        if r["src"] == "offerings" and r.get("teacher_last"):
            want = r["teacher_last"].lower().strip()
            cands = full_by_last.get(want, set())
            if not cands:                                  # two-word last names ("Navarro Rajoy"): match the end of the full name
                wn = " " + norm(want)
                cands = {n for ns in full_by_last.values() for n in ns if (" " + norm(n)).endswith(wn) and len(wn) > 4}
            if not cands:                                  # spelled slightly differently between the two sheets
                cands = {n for k, ns in full_by_last.items() if edit1(k, want) for n in ns}
            if len(cands) == 1:
                r["teacher"] = next(iter(cands))
                if last_name(r["teacher"]) != want:
                    aliases[r["teacher"]].add(r["teacher_last"].strip())
            else:
                r["teacher"] = r["teacher_last"]
        key = (r["semester"], norm(r["course"]), r["block"], last_name(r.get("teacher", "")))
        cur = merged.get(key)
        if cur is None:
            merged[key] = dict(r)
        else:
            for k in ("room", "dept", "home", "teacher"):
                if len(str(r.get(k, ""))) > len(str(cur.get(k, ""))):
                    cur[k] = r[k]
    named = {(k[0], k[1], k[2]) for k, v in merged.items() if v.get("teacher")}
    rows = [v for k, v in merged.items() if v.get("teacher") or (k[0], k[1], k[2]) not in named]
    chunks = []
    file = file_master if master_wb is not None else file_off
    day_of = lambda b: "A/A' days" if b <= 4 else "B/B' days"

    by_teacher = defaultdict(list)
    by_course = defaultdict(list)
    by_block = defaultdict(list)
    for r in rows:
        if r.get("teacher") and " " in r["teacher"]:
            by_teacher[(r["semester"], r["teacher"])].append(r)
        by_course[(r["semester"], norm(r["course"]))].append(r)
        by_block[(r["semester"], r["block"])].append(r)

    for (sem, teacher), rs in by_teacher.items():
        rs.sort(key=lambda x: (x["block"], x["course"]))
        dept = next((x["dept"] for x in rs if x["dept"]), "")
        home = next((x["home"] for x in rs if x["home"]), "")
        lines = [f"- Block {x['block']} ({day_of(x['block'])}): {x['course']}" + (f", room {x['room']}" if x["room"] else "") for x in rs]
        alt = sorted(aliases.get(teacher, set()))
        body = (f"{teacher} teaches ({sem})" + (f" in {dept}" if dept else "") + (f"; assigned room {home}" if home else "") + ":\n"
                + "\n".join(lines))
        if alt:
            body += f"\n(The course offerings sheet spells this last name: {', '.join(alt)}.)"
        chunks.append(mk(f"teacher|{sem}|{teacher}", file, f"{teacher}: courses, blocks and rooms ({sem})", "Teaching assignments", body,
                         {"source": "school_courses", "kind": "teacher", "semester": sem, "teacher": teacher,
                          "teacher_last": last_name(teacher), "aliases": [a.lower() for a in alt], "dept": dept}))
    for (sem, ckey), rs in by_course.items():
        rs.sort(key=lambda x: (x["block"], x.get("teacher", "")))
        name = max((x["course"] for x in rs), key=len)
        lines = []
        for x in rs:
            who = x.get("teacher") or "teacher to be announced"
            lines.append(f"- Block {x['block']} ({day_of(x['block'])}): {who}" + (f", room {x['room']}" if x["room"] else ""))
        teachers = sorted({x["teacher"] for x in rs if x.get("teacher")})
        body = f"{name} ({sem}) is offered in these sections:\n" + "\n".join(lines)
        if teachers:
            body += f"\nTaught by: {', '.join(teachers)}."
        chunks.append(mk(f"course|{sem}|{ckey}", file, f"{name}: who teaches it, blocks and rooms ({sem})", "Course sections", body,
                         {"source": "school_courses", "kind": "course", "semester": sem, "course": name, "course_key": ckey,
                          "teachers": teachers}))
    for (sem, b), rs in by_block.items():
        items = sorted({f"{x['course']}" + (f" ({x['teacher']})" if x.get("teacher") else "") for x in rs})
        for i in range(0, len(items), 40):
            part = items[i:i + 40]
            chunks.append(mk(f"block|{sem}|{b}|{i}", file, f"Courses offered in Block {b} ({sem})", "Offerings by block",
                             f"Block {b} ({day_of(b)}) offerings, {sem}:\n" + "\n".join(f"- {p}" for p in part),
                             {"source": "school_courses", "kind": "block", "semester": sem, "block": b}))
    return chunks


# --------------------------------------------------------------------------- clubs
def club_chunks(wb, file: str) -> list[dict]:
    ws = wb.worksheets[0]
    rows = grid(ws)
    head = [cell(c).lower() for c in rows[0]]
    idx = {n: head.index(n) for n in ("clubs", "description", "leaders", "advisor", "day", "time", "room") if n in head}
    chunks, outline = [], []
    for r in rows[1:]:
        r = r + [None] * 8
        g = lambda k: cell(r[idx[k]]) if k in idx else ""
        name = g("clubs")
        if not name:
            continue
        meets = ", ".join(x for x in (g("day"), g("time")) if x)
        lines = [f"Club: {name} (2026-27)"]
        if g("description"):
            lines.append(f"What it does: {g('description')}")
        if g("leaders"):
            lines.append(f"Student leaders: {g('leaders')}")
        if g("advisor"):
            lines.append(f"Faculty advisor: {g('advisor')}")
        if meets:
            lines.append(f"Meets: {meets}" + (f" in room {g('room')}" if g("room") else ""))
        chunks.append(mk(f"club|{name}", file, f"{name} (club)", "Clubs catalog 2026-27", "\n".join(lines),
                         {"source": "school_clubs", "kind": "club", "club": name}))
        outline.append(f"- {name}" + (f" ({meets})" if meets else ""))
    for i in range(0, len(outline), 30):
        chunks.append(mk(f"clubs-outline|{i}", file, "All 2026-27 clubs and affinity groups", "Clubs catalog 2026-27 (outline)",
                         f"2026-27 clubs and affinity groups, {len(outline)} listed (name and when it meets):\n" + "\n".join(outline[i:i + 30]),
                         {"source": "school_clubs", "kind": "outline"}))
    return chunks


# --------------------------------------------------------------------------- entry point
def find_workbooks() -> dict[str, list[Path]]:
    found: dict[str, list[Path]] = defaultdict(list)
    for folder in (config.DATA_DIR / "live", config.DATA_DIR / "manual_docs"):     # live copies win over manual ones
        for p in sorted(folder.glob("*.xlsx")) if folder.exists() else []:
            if p.name.startswith("~$"):
                continue
            k = kind_of_path(p)
            if k:
                found[k].append(p)
    return found


def build_chunks(log=print) -> list[dict]:
    found = find_workbooks()
    chunks: list[dict] = []
    master = off = None
    mfile = ofile = ""
    if found.get("master"):
        p = found["master"][0]
        master, mfile = openpyxl.load_workbook(str(p), data_only=True), p.stem
        dc, ac = day_chunks(master, mfile), advisor_chunks(master, mfile)
        chunks += dc + ac
        log(f"  master schedule '{p.name}': {len(dc)} day records, {len(ac)} advisor records")
    if found.get("offerings"):
        p = found["offerings"][0]
        off, ofile = openpyxl.load_workbook(str(p), data_only=True), p.stem
    if master is not None or off is not None:
        cc = course_chunks(master, off, mfile, ofile)
        chunks += cc
        kinds = defaultdict(int)
        for c in cc:
            kinds[c["metadata"]["kind"]] += 1
        log(f"  course map: {dict(kinds)} (offerings from {found['offerings'][0].name if found.get('offerings') else 'none'})")
    for p in found.get("clubs", [])[:1]:
        cl = club_chunks(openpyxl.load_workbook(str(p), data_only=True), p.stem)
        chunks += cl
        log(f"  clubs '{p.name}': {len(cl)} records")
    return chunks
