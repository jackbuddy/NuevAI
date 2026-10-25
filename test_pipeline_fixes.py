"""Offline regression tests for the parsing bugs found by scripts/audit_data.py (no network, no database).

    python test_pipeline_fixes.py
"""
import sys
from datetime import date, datetime

import openpyxl

import freshness
from pipeline import calendar_feed, sheets, workbooks

FAILS = []


def check(name, cond, detail=""):
    print(f"{'ok  ' if cond else 'FAIL'} {name}{'' if cond else '  ' + str(detail)}")
    if not cond:
        FAILS.append(name)


# --- freshness -----------------------------------------------------------------------------------------------
T = date(2026, 10, 25)
c = freshness.classify("2025-2026 Course Requests Office Hours\nSign up below", "google_doc", T)
check("freshness: a prior-year title is stale", c["stale"] and c["school_year"] == "2025-26", c)
c = freshness.classify("2026-27 Course Catalog\nIn the 2025-26 school year we added Climate Science", "google_doc", T)
check("freshness: the current year in the title wins", not c["stale"] and c["school_year"] == "2026-27", c)
body = "Student Handbook\n" + "text " * 400 + "\nIn the 2024-25 school year the library was renovated."
check("freshness: a passing mention deep in a page does not make it stale", not freshness.classify(body, "google_doc", T)["stale"])
check("freshness: ...except on the public website, where it marks a year-specific statement",
      freshness.classify(body, "nexus", T, whole_page=True)["stale"])
check("freshness: 'Updated April 28, 2025' in the head is stale", freshness.classify("Scheduling FAQs\nUpdated April 28, 2025\nDear Mavericks", "google_doc", T)["stale"])
check("freshness: a July 2026 notice is not stale", not freshness.classify("Action required\nSubmit by July 21, 2026", "nexus", T)["stale"])
check("freshness: ISO dates are not school years", freshness.school_years_in("on 2026-10-22 and 2026-11-03") == set())
check("freshness: structured sources are current by construction", not freshness.classify("2024-25 anything", "school_day", T)["stale"])
check("freshness: title_key ignores the year", freshness.title_key("2025-26 Course Catalog") == freshness.title_key("2026-27 Course Catalog"))
check("freshness: superseded clubs draft", freshness.classify("Clubs & Affinity Groups Catalog (Under Construction)\nAI & ML Club", "google_sheet", T)["stale"])

# --- iCal: an all-day event whose DTEND is a time on the same day (Daylight Saving Time Ends) ----------------------
ics = """BEGIN:VCALENDAR
BEGIN:VEVENT
UID:a
DTSTART;VALUE=DATE:20261101
DTEND;TZID=America/Los_Angeles:20261101T010000
SUMMARY:Daylight Saving Time Ends
END:VEVENT
BEGIN:VEVENT
UID:b
DTSTART;VALUE=DATE:20261123
DTEND;VALUE=DATE:20261128
SUMMARY:Thanksgiving Break
END:VEVENT
BEGIN:VEVENT
UID:c
DTSTART;TZID=America/Los_Angeles:20261104T112000
DTEND;TZID=America/Los_Angeles:20261104T125000
SUMMARY:Diwali
END:VEVENT
END:VCALENDAR
"""
evs = {e["summary"]: calendar_feed.describe(e) for e in calendar_feed.parse_ics(ics)}
check("ical: all-day event with a timed DTEND stays on its day", evs["Daylight Saving Time Ends"]["end_date"] == date(2026, 11, 1), evs["Daylight Saving Time Ends"])
check("ical: all-day DTEND is exclusive", evs["Thanksgiving Break"]["end_date"] == date(2026, 11, 27), evs["Thanksgiving Break"])
check("ical: timed event keeps its wall time", evs["Diwali"]["time"] == "11:20 AM to 12:50 PM", evs["Diwali"])

# --- team sheets: a date pasted into the Home cell, and the repeated home-games table ---------------------------------
csv = """Date,Home Team,Away Team,Location,Time
Tue 9/8/26,Tue 9/8/26,Nueva,Eastside College Prep,6:00 PM
Thu 9/10/26,Nueva,Crystal Springs,Nueva - San Mateo Campus,5:30 PM
Varsity Home Schedule,,,,
Date,Home Team,Away Team,Location,Time
Thu 9/10/26,Nueva,Crystal Springs,Nueva - San Mateo Campus,5:30 PM
"""
ev = sheets.sports_events({"text": csv, "title": "US Girls Volleyball", "url": "u"}, date(2026, 1, 1), date(2027, 1, 1))
check("sheets: duplicate home-games table is not repeated", len(ev) == 2, [e["summary"] for e in ev])
check("sheets: a date in the Home cell falls back to the location as opponent", "Eastside College Prep" in ev[0]["summary"], ev[0]["summary"])

# --- master schedule day grid: a date cell with text typed into it ------------------------------------------------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Calendar AB + Special Schedule"
ws.append(["Week", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])
ws.append([None, datetime(2027, 2, 1)])
ws.append([6, 22.0, 23.0, "24 Community (90')", 25.0, 26.0])
ws.append([None, "Advisory", None, "All School Meeting", "Advisory", None])
ws.append([None, "A", "B", "A", "A'", "B'"])
ws.append([None, datetime(2027, 3, 1)])
ws.append([None, None, None, None, None, 1.0])          # a later month row, so the grid has more than one block
ws.append([None, None, None, None, None, "No Classes"])
ws.append([None, None, None, None, None, "Spring Break"])
days = workbooks.parse_days(wb)
check("workbooks: a text-typed date cell still yields its day", days.get(date(2027, 2, 24), {}).get("type") == "A", days.get(date(2027, 2, 24)))
check("workbooks: neighbours are unaffected", days.get(date(2027, 2, 25), {}).get("type") == "A'" and days.get(date(2027, 2, 23), {}).get("type") == "B")
wb2 = openpyxl.Workbook()
w2 = wb2.active
w2.title = "Calendar AB + Special Schedule"
w2.append(["Week", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])
w2.append([None, datetime(2027, 6, 1)])
w2.append([None, None, None, 2.0, 3.0, 4.0])
w2.append([None, None, None, "Full-Day Event", "Full-Day Event", "Last Day of School"])
w2.append([None, None, None, "Quest Day", "RenFaire", "Noon Dismissal"])
chunks = {c["metadata"]["event_date"]: c for c in workbooks.day_chunks(wb2, "t")}
check("workbooks: full-day event is not 'No classes'", chunks["2027-06-02"]["metadata"]["kind"] == "full_day_event" and "No classes" not in chunks["2027-06-02"]["content"])
check("workbooks: noon dismissal is not 'No classes'", chunks["2027-06-04"]["metadata"]["kind"] == "short_day" and "No classes" not in chunks["2027-06-04"]["content"],
      chunks["2027-06-04"]["content"][-200:])

# --- courses: Roles wins over the offerings sheet for the same teacher and block; joint classes split -----------------------
m = openpyxl.Workbook()
r = m.active
r.title = "Fall Roles"
r.append(["Fall Semester", None, "A | A' Days", None, None, None, None, "B | B' Days"])
r.append([None, None, "Block 1", "Block 2", "Block 3", "Block 4", None, "Block 5", "Block 6", "Block 7", "Block 8"])
r.append(["Languages"])
r.append(["Laura Mendez-Ortiz", None, None, "Spanish 4\n343B", None, None, None, None, None, "Spanish 3\n1206"])
r.append(["Jehnna Ronan", None, None, None, None, None, None, None, "Bio Teams 2\n304"])
r.append(["Lindy Jensen", None, None, None, None, None, None, None, "Bio Teams 2\n304"])
o = openpyxl.Workbook()
s1 = o.active
s1.title = "Sheet1"
s1.append(["Block 1", "Block 2", "Block 3", "Block 4", None, "Block 5", "Block 6", "Block 7", "Block 8"])
s1.append([None, None, None, None, None, None, "Bio Teams 2\nJensen + Ronan", "Spanish 4\nMendez-Ortiz", None])
ch = workbooks.course_chunks(m, o, "m", "o")
teach = {c["metadata"]["teacher"]: c["content"] for c in ch if c["metadata"].get("kind") == "teacher" and "Fall" in c["metadata"].get("semester", "")}
check("courses: Roles entry kept, double-booked offerings entry dropped", "Block 7 (B/B' days): Spanish 3, room 1206" in teach["Laura Mendez-Ortiz"] and "Spanish 4\n" not in teach["Laura Mendez-Ortiz"].split("Block 7")[-1],
      teach["Laura Mendez-Ortiz"])
check("courses: a joint class is credited to both teachers, no 'Jensen + Ronan' record", set(teach) == {"Laura Mendez-Ortiz", "Jehnna Ronan", "Lindy Jensen"}, set(teach))

print(f"\n{'all passed' if not FAILS else str(len(FAILS)) + ' FAILED: ' + ', '.join(FAILS)}")
sys.exit(1 if FAILS else 0)
